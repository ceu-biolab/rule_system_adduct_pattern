import json

import pytest
from openpyxl import Workbook, load_workbook

from lipid_rule_eval.metrics.evaluator import (
    candidate_sum_key,
    evaluate,
    match_references,
    rank_candidates,
    read_gold_standard,
    sum_composition,
)


def test_sum_composition_combines_resolved_lipid_chains():
    assert sum_composition("PC(16:0/18:1)") == (34, 1)
    assert sum_composition("TG(16:0/18:1/18:2)") == (52, 3)
    assert sum_composition("Cer(d18:1/24:0)") == (42, 1)


def test_candidate_key_prefers_cmm_composition_and_keeps_subclass():
    match = {
        "rule_target": "LPC",
        "compound_name": "PC(16:0/0:0)",
        "cmm_number_carbons": 16,
        "cmm_double_bonds": 0,
    }

    assert candidate_sum_key(match) == "LPC(16:0)"


def _candidate(identifier, score, reference=False):
    return {
        "candidate_id": identifier,
        "group_id": "G1",
        "polarity": "POSITIVE",
        "normalized_score": score,
        "reference_supported": reference,
        "top1": False,
        "average_rank": None,
        "retained": True,
    }


def test_top1_accepts_ties_but_mrr_uses_average_rank():
    candidates = [
        _candidate("reference", 1.0, True),
        _candidate("tied", 1.0),
        _candidate("lower", 0.0),
        _candidate("no-evidence", None),
    ]

    groups = rank_candidates(candidates)

    reference = candidates[0]
    assert reference["top1"] is True
    assert reference["average_rank"] == 1.5
    assert reference["retained"] is True
    assert candidates[1]["top1"] is True
    assert candidates[2]["retained"] is False
    assert candidates[3]["retained"] is True
    assert groups[0]["reference_is_top1"] is True
    assert groups[0]["reference_reciprocal_rank"] == 2.0 / 3.0


def test_reference_match_requires_sum_polarity_adduct_mz_and_rt():
    candidate = {
        "candidate_id": "candidate-1",
        "group_id": "G1",
        "hypothesis_id": "H1",
        "polarity": "POSITIVE",
        "sum_key": "PC(34:1)",
        "items": [
            {
                "mz": 760.5851,
                "rt_seconds": 600.0,
                "adduct_name": "[M+H]+",
            }
        ],
        "reference_supported": False,
        "matched_gold_ids": [],
        "matched_gold_names": [],
        "best_mz_error_da": None,
        "best_rt_error_seconds": None,
        "matched_adduct": None,
    }
    reference = {
        "gold_id": "GS-1",
        "polarity": "POSITIVE",
        "sum_key": "PC(34:1)",
        "sum_name": "PC(34:1)",
        "rt_seconds": 605.9,
        "adduct_masses": {"[M+H]+": 760.5850},
    }

    matches = match_references([candidate], [reference], 20.0, 6.0)

    assert len(matches) == 1
    assert candidate["reference_supported"] is True
    assert candidate["matched_adduct"] == "[M+H]+"
    assert matches[0]["mz_error_ppm"] == pytest.approx(0.0001 / 760.5850 * 1e6)


def _reference_case(reference_mz, observed_mz, observed_rt=600.0):
    candidate = {
        "candidate_id": "C1", "group_id": "G1", "hypothesis_id": "H1",
        "polarity": "POSITIVE", "sum_key": "PC(34:1)",
        "items": [{"mz": observed_mz, "rt_seconds": observed_rt, "adduct_name": "[M+H]+"}],
        "reference_supported": False,
    }
    reference = {
        "gold_id": "GS-1", "polarity": "POSITIVE", "sum_key": "PC(34:1)",
        "sum_name": "PC(34:1)", "rt_seconds": 600.0,
        "adduct_masses": {"[M+H]+": reference_mz},
    }
    return candidate, reference


@pytest.mark.parametrize(
    "reference_mz,observed_mz,expected",
    [(100.0, 100.003, False), (1000.0, 1000.003, True),
     (1000.0, 1000.020, True), (1000.0, 1000.0201, False)],
)
def test_twenty_ppm_scales_with_reference_mz_and_includes_boundary(
    reference_mz, observed_mz, expected
):
    candidate, reference = _reference_case(reference_mz, observed_mz)
    matches = match_references([candidate], [reference], 20.0, 6.0)

    assert bool(matches) is expected
    assert candidate["reference_supported"] is expected


@pytest.mark.parametrize("mismatch", ["polarity", "sum", "adduct", "rt"])
def test_ppm_matching_preserves_polarity_sum_adduct_and_rt_requirements(mismatch):
    candidate, reference = _reference_case(760.0, 760.005)
    if mismatch == "polarity":
        candidate["polarity"] = "NEGATIVE"
    elif mismatch == "sum":
        candidate["sum_key"] = "PC(36:2)"
    elif mismatch == "adduct":
        candidate["items"][0]["adduct_name"] = "[M+Na]+"
    else:
        candidate["items"][0]["rt_seconds"] = 606.01

    assert match_references([candidate], [reference], 20.0, 6.0) == []


@pytest.mark.parametrize("ppm", [0.0, -1.0, float("nan"), float("inf")])
def test_rejects_invalid_ppm_tolerance(ppm):
    with pytest.raises(ValueError, match="finite and positive"):
        match_references([], [], ppm, 6.0)


@pytest.fixture
def fedorova_gold(tmp_path):
    path = tmp_path / "fedorova.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "S1"
    sheet.append(["Table S1"])
    sheet.append([
        "Class", "Shorthand notation: lipid species level", "Molecular species",
        "LipidLynxX species", "LipidLynxX molecular species", "Formula_Neutral",
        "Exact_mass", "[M+H]+", "[M-H2O+H]+", "[M+NH4]+", "[M+HCOO]-",
        "[M-H]-", "HILIC", "C18 RPC", "C30 RPC", "C30* RPC",
    ])
    sheet.append([None])
    sheet.append(["carnitines"])
    sheet.append([
        "PC", "PC 34:1", "PC 16:0_18:1", "PC(34:1)", "PC(16:0_18:1)",
        "C42H82NO8P", 759.577, "[M+H]+", None, None, "[M+HCOO]-", None,
        None, 10.0, 12.0, None,
    ])
    sheet.append([
        "Cer", "Cer 34:1;O2", "Cer 18:1;O2/16:0", "Cer(34:1;O2)",
        "Cer(18:1;O2_16:0)", "C34H67NO3", 537.5121, None, "[M-H2O+H]+",
        None, None, None, None, 11.0, None, None,
    ])
    sheet.append([
        "CE", "CE  18:1", "CE  18:1", "CE  18:1", "CE  18:1",
        "C45H78O2", 650.6002, None, None, "[M+NH4]+", None, None,
        None, None, 24.99, None,
    ])
    workbook.save(path)
    return path


def test_fedorova_requires_explicit_chromatography_method(fedorova_gold):
    with pytest.raises(ValueError, match="requires --reference-method"):
        read_gold_standard(fedorova_gold)


def test_fedorova_normalizes_method_rts_polarities_and_dehydrated_adduct(fedorova_gold):
    references = read_gold_standard(fedorova_gold, reference_method="C18 RPC")

    assert len(references) == 3
    assert {item["polarity"] for item in references} == {"POSITIVE", "NEGATIVE"}
    pc_pos = next(item for item in references if item["sum_key"] == "PC(34:1)" and item["polarity"] == "POSITIVE")
    assert pc_pos["rt_seconds"] == 600.0
    assert pc_pos["adduct_masses"]["[M+H]+"] == pytest.approx(760.584276)
    assert pc_pos["source_row"] == 5
    cer = next(item for item in references if item["sum_key"] == "CER(34:1)")
    assert cer["adduct_masses"]["[M+H-H2O]+"] == pytest.approx(520.5089)
    c30_refs = read_gold_standard(fedorova_gold, reference_method="C30 RPC")
    assert next(item for item in c30_refs if item["sum_key"] == "PC(34:1)")["rt_seconds"] == 720.0
    assert next(item for item in c30_refs if item["sum_key"] == "CE(18:1)")["rt_seconds"] == pytest.approx(1499.4)
    assert not any(item["sum_key"] == "CER(34:1)" for item in c30_refs)


def test_fedorova_evaluation_writes_ppm_and_method_to_workbook(fedorova_gold, tmp_path):
    path = tmp_path / "scored.json"
    path.write_text(json.dumps({"groups": [{
        "group_id": "G1", "polarity": "POSITIVE", "source_files": ["C18.mzML"],
        "hypotheses": [{
            "hypothesis_id": "H1", "neutral_mass": 759.577,
            "items": [{"mz": 760.590, "rt_seconds": 600.0, "adduct_name": "[M+H]+"}],
            "compound_matches": [{"compound_id": 1, "compound_name": "PC(16:0/18:1)", "rule_target": "PC"}],
            "rule_scores": [{"rule_target": "PC", "execution_status": "SCORED", "score": 1,
                             "normalized_score": 1.0, "has_rule_evidence": True,
                             "evidence_status": "SUPPORTING", "applied_presence": 1, "applied_intensity": 0}],
        }],
    }]}))
    output = tmp_path / "metrics.xlsx"

    report = evaluate(path, fedorova_gold, output, reference_method="C18 RPC")

    assert report["reference_match_details"] == 1
    workbook = load_workbook(output, read_only=True, data_only=True)
    parameters = dict(workbook["Parameters"].iter_rows(min_row=2, values_only=True))
    assert parameters["m/z tolerance (ppm)"] == 20.0
    assert parameters["reference chromatography method"] == "C18 RPC"
    headers = next(workbook["Reference matches"].iter_rows(values_only=True))
    assert "mz_error_ppm" in headers
    workbook.close()
