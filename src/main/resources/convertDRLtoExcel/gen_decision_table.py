#!/usr/bin/env python3
"""
Parse all .drl files and generate six Drools decision table Excel files:
  positive_presence.xlsx, positive_intensityGT.xlsx, positive_intensityLT.xlsx
  negative_presence.xlsx, negative_intensityGT.xlsx, negative_intensityLT.xlsx
"""

import re
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_HERE    = Path(__file__).resolve().parent
DRL_BASE = _HERE.parent / "rules"
OUT_DIR  = _HERE.parent / "rules"

# ── Regex patterns ──────────────────────────────────────────────────────────
RULE_RE      = re.compile(r'rule\s+"([^"]+)"\s*\nwhen(.*?)then(.*?)end', re.DOTALL)
PRESENCE_RE  = re.compile(r'\$adduct\s*:\s*ResultItem\s*\(\s*adductName\s*==\s*"([^"]+)"\s*\)')
ABSENCE_RE   = re.compile(r'not\s+ResultItem\s*\(\s*adductName\s*==\s*"([^"]+)"\s*\)')
ADDUCT1_RE   = re.compile(r'\$adduct1\s*:\s*ResultItem\s*\(\s*adductName\s*==\s*"([^"]+)"\s*\)')
ADDUCT2_RE   = re.compile(r'\$adduct2\s*:\s*ResultItem\s*\(\s*adductName\s*==\s*"([^"]+)"\s*\)')
EVAL_GT_RE   = re.compile(r'eval\(\s*\$adduct1\.getIntensity\(\)\s*>\s*\$adduct2\.getIntensity\(\)\s*\)')
PHASE_RE     = re.compile(r'eval\(\$phase\d+\s*==\s*MobilePhases\.(\w+)\)')
SCORE_RE     = re.compile(r'lipid\.setScore\((-?\d+)\)')
DESCR_OK_RE  = re.compile(r'lipid\.setDescrCorrect\("([^"]+)"\)')
DESCR_BAD_RE = re.compile(r'lipid\.setDescrIncorrect\("([^"]+)"\)')


# ── DRL parsing ─────────────────────────────────────────────────────────────
def parse_drl(path: Path, polarity: str) -> list:
    text = path.read_text(encoding='utf-8')
    rules = []
    for m in RULE_RE.finditer(text):
        name, when_blk, then_blk = m.group(1), m.group(2), m.group(3)
        if 'Dummy' in name:
            continue

        a1   = ADDUCT1_RE.search(when_blk)
        a2   = ADDUCT2_RE.search(when_blk)
        pres = PRESENCE_RE.search(when_blk)
        abs_ = ABSENCE_RE.search(when_blk)
        phases = PHASE_RE.findall(when_blk)

        sm = SCORE_RE.search(then_blk)
        score = int(sm.group(1)) if sm else 0
        dm  = DESCR_OK_RE.search(then_blk)
        dc  = dm.group(1) if dm else ''
        dm2 = DESCR_BAD_RE.search(then_blk)
        di  = dm2.group(1) if dm2 else ''

        if a1 and a2:
            gt = bool(EVAL_GT_RE.search(when_blk))
            rules.append({'name': name,
                          'polarity': polarity,
                          'type': 'intensity_gt' if gt else 'intensity_lt',
                          'adduct1': a1.group(1), 'adduct2': a2.group(1),
                          'phases': phases, 'score': score,
                          'descr_correct': dc, 'descr_incorrect': di})
        elif pres:
            rules.append({'name': name, 'polarity': polarity, 'type': 'presence',
                          'adduct': pres.group(1),
                          'phases': phases, 'score': score,
                          'descr_correct': dc, 'descr_incorrect': ''})
        elif abs_:
            rules.append({'name': name, 'polarity': polarity, 'type': 'absence',
                          'adduct': abs_.group(1),
                          'phases': phases, 'score': score,
                          'descr_correct': '', 'descr_incorrect': di})
    return rules


def parse_all(base: Path) -> list:
    rules = []
    for f in sorted(base.rglob('*.drl')):
        if 'userFiles' in str(f):
            continue
        parts = f.parts
        if 'positive' in parts:
            polarity = 'positive'
        elif 'negative' in parts:
            polarity = 'negative'
        else:
            polarity = 'unknown'
        rules.extend(parse_drl(f, polarity))
    return rules


# ── Styles ───────────────────────────────────────────────────────────────────
HDR_FILL   = PatternFill('solid', fgColor='4472C4')
COND_FILL  = PatternFill('solid', fgColor='E2EFDA')
ACT_FILL   = PatternFill('solid', fgColor='FFF2CC')
LBL_FILL   = PatternFill('solid', fgColor='D9E1F2')
WHITE_BOLD = Font(bold=True, color='FFFFFF')
BOLD       = Font(bold=True)


# ── Excel helpers ─────────────────────────────────────────────────────────────
def write_metadata(ws, start_row: int) -> int:
    r = start_row
    entries = [
        ('RuleSet',   'ceu.biolab.cmm.rulePuntuation.rules'),
        ('Import',    'ceu.biolab.cmm.shared.dto.FeatureAnnotation.AnnotatedFeature,'
                      'ceu.biolab.cmm.shared.dto.FeatureAnnotation.ResultItem,'
                      'ceu.biolab.cmm.shared.domain.MobilePhases,'
                      'java.util.List'),
        ('Variables', 'AnnotatedFeature lipid'),
        ('Variables', 'List mobilePhases'),
        ('Variables', 'String sampleType'),
    ]
    for key, val in entries:
        ws.cell(r, 2, key).font = BOLD
        ws.cell(r, 3, val)
        r += 1
    return r + 1


def write_table(ws, start_row: int, table_name: str,
                types: list, declarations: list, snippets: list,
                labels: list, data: list) -> int:
    r = start_row

    c = ws.cell(r, 2, f'RuleTable {table_name}')
    c.fill = HDR_FILL; c.font = WHITE_BOLD
    r += 1

    ws.cell(r, 1, 'DESCRIPTION').font = BOLD
    for ci, t in enumerate(types, start=2):
        cell = ws.cell(r, ci, t)
        cell.font = BOLD
        if t == 'CONDITION': cell.fill = COND_FILL
        elif t == 'ACTION':  cell.fill = ACT_FILL
    r += 1

    for ci, decl in enumerate(declarations, start=2):
        if decl: ws.cell(r, ci, decl)
    r += 1

    for ci, snip in enumerate(snippets, start=2):
        if snip: ws.cell(r, ci, snip)
    r += 1

    for ci, lbl in enumerate(labels, start=2):
        if lbl:
            cell = ws.cell(r, ci, lbl)
            cell.fill = LBL_FILL
    r += 1

    for row_vals in data:
        for ci, val in enumerate(row_vals, start=2):
            if val is not None and val != '':
                ws.cell(r, ci, val)
        r += 1

    return r + 1


# ── Table builders ────────────────────────────────────────────────────────────
def build_presence_table(rules: list, polarity: str):
    subset = [r for r in rules if r['polarity'] == polarity and r['type'] in ('presence', 'absence')]
    table_name = f'{polarity.capitalize()}PresenceRules'

    cols = [
        ('NAME',      '',                                  '',                                                        'Rule Name'),
        ('CONDITION', 'ResultItem',                        'adductName == "$1"',                                      'Adduct Present'),
        ('CONDITION', 'not ResultItem',                    'adductName == "$1"',                                      'Adduct Absent'),
        ('CONDITION', 'MobilePhases from mobilePhases',   'this == MobilePhases.$1',                                 'Phase 1'),
        ('CONDITION', 'MobilePhases from mobilePhases',   'this == MobilePhases.$1',                                 'Phase 2'),
        ('CONDITION', 'MobilePhases from mobilePhases',   'this == MobilePhases.$1',                                 'Phase 3'),
        ('ACTION',    '',                                  'lipid.setScore($1);',                                     'Score'),
        ('ACTION',    '',                                  'lipid.setDescrCorrect("$1");',                            'Descr Correct'),
        ('ACTION',    '',                                  'lipid.setDescrIncorrect("$1");',                          'Descr Incorrect'),
        ('ACTION',    '',                                  'lipid.setAppliedPresence(lipid.getAppliedPresence() + 1);', 'Applied Presence'),
    ]
    types = [c[0] for c in cols]; declarations = [c[1] for c in cols]
    snippets = [c[2] for c in cols]; labels = [c[3] for c in cols]

    data = []
    for rule in subset:
        ph = (rule['phases'] + ['', '', ''])[:3]
        if rule['type'] == 'presence':
            data.append([rule['name'], rule['adduct'], '',
                         ph[0], ph[1], ph[2],
                         rule['score'], rule['descr_correct'], '', 'x'])
        else:
            data.append([rule['name'], '', rule['adduct'],
                         ph[0], ph[1], ph[2],
                         rule['score'], '', rule['descr_incorrect'], 'x'])

    return table_name, types, declarations, snippets, labels, data


def build_intensity_table(rules: list, polarity: str, is_gt: bool):
    rule_type  = 'intensity_gt' if is_gt else 'intensity_lt'
    subset     = [r for r in rules if r['polarity'] == polarity and r['type'] == rule_type]
    tag        = 'GT' if is_gt else 'LT'
    table_name = f'{polarity.capitalize()}IntensityRules{tag}'
    eval_snip  = ('eval($a1.getIntensity() > $a2.getIntensity())'
                  if is_gt else
                  'eval($a1.getIntensity() < $a2.getIntensity())')

    cols = [
        ('NAME',      '',                                  '',                                                          'Rule Name'),
        ('CONDITION', '$a1: ResultItem',                   'adductName == "$1"',                                        'Higher adduct ($a1)' if is_gt else 'Lower adduct ($a1)'),
        ('CONDITION', '$a2: ResultItem',                   'adductName == "$1"',                                        'Other adduct ($a2)'),
        ('CONDITION', '',                                  eval_snip,                                                   'Intensity eval'),
        ('CONDITION', 'MobilePhases from mobilePhases',   'this == MobilePhases.$1',                                   'Phase 1'),
        ('CONDITION', 'MobilePhases from mobilePhases',   'this == MobilePhases.$1',                                   'Phase 2'),
        ('CONDITION', 'MobilePhases from mobilePhases',   'this == MobilePhases.$1',                                   'Phase 3'),
        ('ACTION',    '',                                  'lipid.setScore($1);',                                       'Score'),
        ('ACTION',    '',                                  'lipid.setDescrCorrect("$1");',                              'Descr Correct'),
        ('ACTION',    '',                                  'lipid.setDescrIncorrect("$1");',                            'Descr Incorrect'),
        ('ACTION',    '',                                  'lipid.setAppliedIntensity(lipid.getAppliedIntensity() + 1);', 'Applied Intensity'),
    ]
    types = [c[0] for c in cols]; declarations = [c[1] for c in cols]
    snippets = [c[2] for c in cols]; labels = [c[3] for c in cols]

    data = []
    for rule in subset:
        ph = (rule['phases'] + ['', '', ''])[:3]
        data.append([rule['name'],
                     rule['adduct1'], rule['adduct2'], 'x',
                     ph[0], ph[1], ph[2],
                     rule['score'], rule['descr_correct'], rule['descr_incorrect'], 'x'])

    return table_name, types, declarations, snippets, labels, data


# ── Per-file generation ───────────────────────────────────────────────────────
def generate_file(out: Path, table_name: str,
                  types, declarations, snippets, labels, data):
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = table_name[:31]

    row = write_metadata(ws, 1)
    write_table(ws, row, table_name, types, declarations, snippets, labels, data)

    for col_cells in ws.columns:
        max_w = max((len(str(c.value)) for c in col_cells if c.value), default=8)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max_w + 2, 60)

    wb.save(out)
    print(f"Saved: {out} ({len(data)} rules)")


# ── Main ──────────────────────────────────────────────────────────────────────
FILES = [
    ('positive', 'presence',    lambda r: build_presence_table(r, 'positive')),
    ('positive', 'intensityGT', lambda r: build_intensity_table(r, 'positive', True)),
    ('positive', 'intensityLT', lambda r: build_intensity_table(r, 'positive', False)),
    ('negative', 'presence',    lambda r: build_presence_table(r, 'negative')),
    ('negative', 'intensityGT', lambda r: build_intensity_table(r, 'negative', True)),
    ('negative', 'intensityLT', lambda r: build_intensity_table(r, 'negative', False)),
]

if __name__ == '__main__':
    rules = parse_all(DRL_BASE)
    total = 0
    for polarity, kind, builder in FILES:
        out = OUT_DIR / f'{polarity}_{kind}.xlsx'
        table_name, types, decls, snips, labels, data = builder(rules)
        generate_file(out, table_name, types, decls, snips, labels, data)
        total += len(data)
    print(f"\nTotal rules across all files: {total}")
